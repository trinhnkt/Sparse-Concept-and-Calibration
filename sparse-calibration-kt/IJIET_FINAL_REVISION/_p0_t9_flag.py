#!/usr/bin/env python3
"""Table 9 XES very-sparse flag: printed N=114 → L. No ECE/N edits."""
from __future__ import annotations

import shutil
import sys
from pathlib import Path

from docx import Document

HERE = Path(__file__).resolve().parent
FULL = HERE / "manuscript" / "Reproducible Sparse-Concept and Calibration Diagnostics for Knowledge Tracing.docx"
BAK = HERE / "manuscript" / "Reproducible Sparse-Concept and Calibration Diagnostics for Knowledge Tracing.docx.bak_pre_p0_t9_flag"


def main() -> None:
    sys.stdout.reconfigure(encoding="utf-8")
    if not BAK.exists():
        shutil.copy2(FULL, BAK)
    d = Document(str(FULL))
    hits = 0
    for tbl in d.tables:
        rows = [[c.text.strip() for c in row.cells] for row in tbl.rows]
        if not any("XES3G5M" in r[0] and "very-sparse" in r[1] for r in rows if len(r) >= 2):
            continue
        cell = tbl.rows[4].cells[3]
        if cell.text.strip() != "I/L":
            raise SystemExit(f"flag is {cell.text!r}, expected I/L")
        n_cell = tbl.rows[4].cells[2].text.strip()
        tkt = tbl.rows[4].cells[4].text.strip()
        dkt = tbl.rows[4].cells[5].text.strip()
        assist = tbl.rows[2].cells[4].text.strip()
        if n_cell != "114" or tkt != "0.184" or dkt != "0.173":
            raise SystemExit(f"XES lock drifted N={n_cell} T-KT={tkt} DKT={dkt}")
        if assist != "0.245":
            raise SystemExit("Assist very-sparse T-KT ECE drifted")
        runs = cell.paragraphs[0].runs
        if not runs:
            raise SystemExit("empty flag runs")
        runs[0].text = "L"
        for r in runs[1:]:
            r.text = ""
        hits += 1
        break
    if hits != 1:
        raise SystemExit(f"hits={hits}")
    found = False
    for tbl in d.tables:
        rows = [[c.text.strip() for c in row.cells] for row in tbl.rows]
        if len(rows) < 5 or len(rows[0]) < 6:
            continue
        if rows[4][:3] != ["XES3G5M", "very-sparse", "114"]:
            continue
        if rows[4][3] != "L" or rows[4][4] != "0.184" or rows[4][5] != "0.173":
            raise SystemExit(f"XES row after patch {rows[4]}")
        found = True
    if not found:
        raise SystemExit("Table 9 XES row missing after patch")
    full = "\n".join(p.text for p in d.paragraphs)
    if "0.1136" not in full or "0.2280" not in full:
        raise SystemExit("T-KT ECE lock missing")
    d.save(str(FULL))
    print("t9 flag -> L")


if __name__ == "__main__":
    main()

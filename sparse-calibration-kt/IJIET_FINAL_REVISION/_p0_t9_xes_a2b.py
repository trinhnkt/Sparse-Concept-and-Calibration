#!/usr/bin/env python3
"""Align Table 9 XES very-sparse to masked a2b. No T-KT ECE/FAR lock edits."""
from __future__ import annotations

import shutil
import sys
from pathlib import Path

import win32com.client as win32
from docx import Document

HERE = Path(__file__).resolve().parent
FULL = HERE / "manuscript" / "Reproducible Sparse-Concept and Calibration Diagnostics for Knowledge Tracing.docx"
BAK = HERE / "manuscript" / "Reproducible Sparse-Concept and Calibration Diagnostics for Knowledge Tracing.docx.bak_pre_t9_xes_a2b"
WD_SAVE = -1


def txt(cell) -> str:
    return cell.Range.Text.replace("\r", " ").replace("\x07", "").strip()


def set_para_text(p, new: str) -> None:
    if not p.runs:
        p.add_run(new)
        return
    p.runs[0].text = new
    for r in p.runs[1:]:
        r.text = ""


def patch_table() -> None:
    word = win32.DispatchEx("Word.Application")
    word.Visible = False
    word.DisplayAlerts = 0
    doc = word.Documents.Open(str(FULL))
    try:
        tbl = None
        for t in range(1, doc.Tables.Count + 1):
            prev = doc.Range(max(0, doc.Tables(t).Range.Start - 280), doc.Tables(t).Range.Start)
            if "Table 9. Cold-start" in prev.Text:
                tbl = doc.Tables(t)
                break
        if tbl is None:
            raise SystemExit("Table 9 not found")
        if "XES" not in txt(tbl.Cell(5, 1)) or "very-sparse" not in txt(tbl.Cell(5, 2)):
            raise SystemExit(f"unexpected T9 r5 {txt(tbl.Cell(5, 1))} {txt(tbl.Cell(5, 2))}")
        if txt(tbl.Cell(5, 3)) != "112":
            raise SystemExit(f"T9 N is {txt(tbl.Cell(5, 3))}, expected 112")
        if txt(tbl.Cell(5, 5)) != "0.183":
            raise SystemExit(f"T9 T-KT ECE is {txt(tbl.Cell(5, 5))}")
        if txt(tbl.Cell(5, 6)) != "0.184":
            raise SystemExit(f"T9 DKT ECE is {txt(tbl.Cell(5, 6))}")
        tbl.Cell(5, 3).Range.Text = "114"
        tbl.Cell(5, 5).Range.Text = "0.184"
        tbl.Cell(5, 6).Range.Text = "0.173"
        tbl.Range.Font.SmallCaps = False
        tbl.Range.Font.AllCaps = False
        if txt(tbl.Cell(3, 5)) != "0.245":
            raise SystemExit("Assist very-sparse T-KT ECE drifted")
        doc.Save()
        print("t9 xes a2b cells")
    finally:
        try:
            doc.Close(WD_SAVE)
        except Exception:
            pass
        word.Quit()


def patch_prose() -> None:
    d = Document(str(FULL))
    n = 0
    old = "XES3G5M very-sparse support is larger (N≈112) but still not a high-N finding."
    new = "XES3G5M very-sparse support is larger (N≈114) but still not a high-N finding."
    for p in d.paragraphs:
        if old not in p.text:
            continue
        set_para_text(p, p.text.replace(old, new, 1))
        n += 1
        break
    if n != 1:
        raise SystemExit(f"prose hits={n}")
    full = "\n".join(p.text for p in d.paragraphs)
    if "N≈114" not in full:
        raise SystemExit("N≈114 missing")
    if "N≈112" in full:
        raise SystemExit("N≈112 still in prose")
    if "0.1136" not in full or "0.2280" not in full:
        raise SystemExit("T-KT ECE lock missing")
    if "TSCDA" in full:
        raise SystemExit("TSCDA reappeared")
    d.save(str(FULL))
    print("prose N≈114")


def main() -> None:
    sys.stdout.reconfigure(encoding="utf-8")
    if not BAK.exists():
        shutil.copy2(FULL, BAK)
    patch_table()
    patch_prose()


if __name__ == "__main__":
    main()

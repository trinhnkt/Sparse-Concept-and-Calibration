#!/usr/bin/env python3
"""B: S7 pointer in IV.B. C: drop Table 9 small-caps; tighten Limited N / FAR CI."""
from __future__ import annotations

import shutil
import sys
from pathlib import Path

import win32com.client as win32
from docx import Document

HERE = Path(__file__).resolve().parent
FULL = HERE / "manuscript" / "Reproducible Sparse-Concept and Calibration Diagnostics for Knowledge Tracing.docx"
BAK = HERE / "manuscript" / "Reproducible Sparse-Concept and Calibration Diagnostics for Knowledge Tracing.docx.bak_pre_p0_bc"
WD_SAVE = -1


def set_para_text(p, new: str) -> None:
    if not p.runs:
        p.add_run(new)
        return
    p.runs[0].text = new
    for r in p.runs[1:]:
        r.text = ""


def fix_table9_caps() -> None:
    word = win32.DispatchEx("Word.Application")
    word.Visible = False
    word.DisplayAlerts = 0
    doc = word.Documents.Open(str(FULL))
    try:
        tbl = None
        for t in range(1, doc.Tables.Count + 1):
            prev = doc.Range(max(0, doc.Tables(t).Range.Start - 220), doc.Tables(t).Range.Start)
            if "Table 9. Cold-start" in prev.Text:
                tbl = doc.Tables(t)
                break
        if tbl is None:
            raise SystemExit("Table 9 not found")
        tbl.Range.Font.SmallCaps = False
        tbl.Range.Font.AllCaps = False
        if "0.245" not in tbl.Cell(3, 5).Range.Text:
            raise SystemExit("Table 9 T-KT very-sparse ECE missing")
        doc.Save()
        print("t9 smallcaps off")
    finally:
        try:
            doc.Close(WD_SAVE)
        except Exception:
            pass
        word.Quit()


def patch_prose() -> None:
    pairs = [
        (
            "and the T-KT cells stay locked.",
            "and the T-KT cells stay locked. Two alternative train-only cut grids "
            "on the same frozen T-KT scores leave that dense-to-sparse rise positive "
            "(Supplementary Table S7).",
        ),
        (
            "A KC-clustered bootstrap on seed 42 yields a 95% interval [0.006, 0.138] for T-KT ΔFAR.",
            "A KC-clustered bootstrap on seed 42 yields a 95% interval [0.006, 0.138] "
            "for T-KT ΔFAR. That interval is wide, and the sparse ECE cell is Limited "
            "(N=415); neither is a high-N finding.",
        ),
    ]
    d = Document(str(FULL))
    hits = 0
    for old, new in pairs:
        n = 0
        for p in d.paragraphs:
            if old not in p.text:
                continue
            set_para_text(p, p.text.replace(old, new, 1))
            n += 1
            hits += 1
        if n != 1:
            raise SystemExit(f"prose hits={n} for {old[:70]!r}")
    full = "\n".join(p.text for p in d.paragraphs)
    if "Supplementary Table S7" not in full:
        raise SystemExit("S7 pointer missing")
    if "neither is a high-N finding" not in full:
        raise SystemExit("Limited/CI tighten missing")
    if "0.1136" not in full or "0.2280" not in full:
        raise SystemExit("T-KT ECE lock missing")
    if "[0.006, 0.138]" not in full:
        raise SystemExit("FAR CI lock missing")
    d.save(str(FULL))
    print(f"prose patched n={hits}")


def main() -> None:
    sys.stdout.reconfigure(encoding="utf-8")
    if not BAK.exists():
        shutil.copy2(FULL, BAK)
    fix_table9_caps()
    patch_prose()


if __name__ == "__main__":
    main()

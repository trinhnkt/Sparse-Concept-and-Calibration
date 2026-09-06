#!/usr/bin/env python3
"""Point Table 3 to Supplementary Table S10. No Table 3 cell edits."""
from __future__ import annotations

import shutil
import sys
from pathlib import Path

from docx import Document

HERE = Path(__file__).resolve().parent
FULL = HERE / "manuscript" / "Reproducible Sparse-Concept and Calibration Diagnostics for Knowledge Tracing.docx"
BAK = HERE / "manuscript" / "Reproducible Sparse-Concept and Calibration Diagnostics for Knowledge Tracing.docx.bak_pre_p0_s10"

OLD = "Fig. 1 summarizes that pipeline."
NEW = (
    "Fig. 1 summarizes that pipeline. Table 3 presents a compact fold-0 "
    "audit; full fold-by-dataset audit records are provided in "
    "Supplementary Table S10 and the review artifact."
)


def set_para_text(p, new: str) -> None:
    if not p.runs:
        p.add_run(new)
        return
    p.runs[0].text = new
    for r in p.runs[1:]:
        r.text = ""


def main() -> None:
    sys.stdout.reconfigure(encoding="utf-8")
    if not BAK.exists():
        shutil.copy2(FULL, BAK)
    d = Document(str(FULL))
    n = 0
    for p in d.paragraphs:
        if OLD not in p.text or "Table 3 records a seven-channel" not in p.text:
            continue
        if "Supplementary Table S10" in p.text:
            print("S10 pointer already present")
            return
        set_para_text(p, p.text.replace(OLD, NEW, 1))
        n += 1
    if n != 1:
        raise SystemExit(f"hits={n}")
    full = "\n".join(p.text for p in d.paragraphs)
    if "Supplementary Table S10" not in full:
        raise SystemExit("S10 pointer missing")
    if "0.1136" not in full or "0.2280" not in full:
        raise SystemExit("ECE lock missing")
    d.save(str(FULL))
    print("s10 pointer patched")


if __name__ == "__main__":
    main()

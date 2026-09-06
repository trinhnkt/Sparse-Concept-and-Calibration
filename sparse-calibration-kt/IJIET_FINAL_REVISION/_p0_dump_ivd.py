#!/usr/bin/env python3
"""Dump IV.D / Table 7–9 context from the living named Word."""
from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

HERE = Path(__file__).resolve().parent
FULL = HERE / "manuscript" / "Reproducible Sparse-Concept and Calibration Diagnostics for Knowledge Tracing.docx"
OUT = HERE / "audit" / "_p0_ivd_dump.txt"


def cell_text(cell) -> str:
    return " ".join(cell.text.split())


def main() -> None:
    sys.stdout.reconfigure(encoding="utf-8")
    d = Document(str(FULL))
    lines: list[str] = []
    lines.append(f"paragraphs={len(d.paragraphs)} tables={len(d.tables)}")
    for i, tbl in enumerate(d.tables, 1):
        rows, cols = len(tbl.rows), len(tbl.columns)
        h = cell_text(tbl.cell(0, 0))[:40]
        lines.append(f"T{i} {rows}x{cols} h={h!r}")
    flag = False
    for i, p in enumerate(d.paragraphs):
        t = p.text.strip()
        if not t:
            continue
        if (
            t.startswith("D. Dataset-dependent")
            or t.startswith("E. Cold-start")
            or t.startswith("Table 7")
            or t.startswith("Table 8")
            or t.startswith("Table 9")
            or "controlled sparsif" in t.lower()
            or "The three analyses" in t
            or "between-KC" in t
            or "Within-KC" in t
            or t.startswith("C. Why datasets")
        ):
            flag = True
        if t.startswith("V. DISCUSSION") or t.startswith("V. D"):
            flag = True
        if flag:
            lines.append(f"p{i} {t[:500]}")
        if t.startswith("VI. CONCLUSION") or t.startswith("VI. C"):
            break
    OUT.write_text("\n".join(lines), encoding="utf-8")
    print("wrote", OUT, "n", len(lines))


if __name__ == "__main__":
    main()

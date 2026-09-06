#!/usr/bin/env python3
"""P0 Pha 1–2: rewrite Abstract/contributions/II.B/GKT/TSCDA/Limitations on named+blind."""
from __future__ import annotations

import sys
from pathlib import Path

from docx import Document

HERE = Path(__file__).resolve().parent

REPLACEMENTS = [
    (
        "Abstract—Knowledge Tracing (KT) systems that skip, remediate, or advance practice typically consume a predicted probability rather than an area-under-the-curve (AUC) score, so population ranking can look acceptable while probabilities are poorly calibrated on rarely practiced knowledge components (KCs). This paper presents TSCDA (Train-only Sparse-Concept Decision Audit), an evaluation instrument—not a new KT architecture—on three public logs (ASSISTments 2012, Junyi Academy, and XES3G5M) with IRT, DKT, and a local Transformer KT baseline (T-KT; not SimpleKT [4]). Lower KC training frequency does not universally degrade discrimination, but calibration can become less reliable in some sparse-concept regimes. On ASSISTments 2012, T-KT expected calibration error (ECE) rises from 0.114 on dense KCs to 0.228 on sparse KCs (Limited occupancy, N≈415); Junyi’s exercise-level KC tagging yields no learner-based sparse stratum (estimability, not a missing table), and XES3G5M T-KT ECE is essentially flat. A locked global threshold at τ=0.7 raises the T-KT false-advance rate (incorrect-response rate among advance decisions) from 0.196 (dense) to 0.268 (sparse) on one ASSISTments fold under Limited occupancy (Nadvance=235; seed-42 ΔFAR=0.072). The primary check is 4/4 unique learner partitions (partition-level mean ΔFAR 0.056, range 0.015–0.087); the five training runs that underlie those partitions have mean 0.047 (sd 0.033) because seeds 2025 and 2026 share one split. The seed-42 95% CI is [0.006, 0.138]. This is a simulated decision gate, not a classroom intervention. Probability-threshold decisions should be validated by KC-frequency stratum rather than on population AUC or ECE alone.",
        "Abstract—Knowledge Tracing (KT) systems that skip, remediate, or advance practice typically consume a predicted probability rather than an area-under-the-curve (AUC) score, so population ranking can look acceptable while probabilities are poorly calibrated on rarely practiced knowledge components (KCs). We specify a train-only sparse-concept evaluation protocol that combines learner-based and temporal views, explicit cold-start definitions, occupancy-aware reporting, 15-bin expected calibration error (ECE), Brier score with the Murphy decomposition (uncertainty − resolution + reliability), reliability diagrams, and L1–L7 leakage control. The protocol is applied to three public logs (ASSISTments 2012, Junyi Academy, and XES3G5M) with IRT, DKT, and a local Transformer KT baseline (T-KT; not SimpleKT [4]); official SimpleKT is scored on ASSISTments AUC/ECE only. Lower KC training frequency does not universally degrade discrimination: Junyi’s exercise-level operational identifier (ucid) yields no learner-based sparse stratum, and XES3G5M T-KT ECE is essentially flat (0.1176, 0.1129, 0.1254), whereas ASSISTments 2012 T-KT ECE rises from 0.114 on dense KCs to 0.228 on sparse KCs (Limited occupancy, N≈415). A frozen review artifact and a one-command rebuild of the diagnostic tables from frozen summaries accompany the protocol; a locked simulated gate at τ=0.7 is reported only as a decision-error probe (Supplementary Tables S3–S6), not as a primary contribution. This is not a classroom intervention.",
    ),
    (
        "Contributions are conservative: (i) TSCDA, a train-only sparse-concept decision audit—frequency strata with a strict cold-start bin, occupancy-gated ECE/Brier, and a locked-τ FAR—so “sparse” cannot leak test-fold counts; (ii) an estimability result: whether that audit has a sparse stratum depends on how the platform defines a KC (skill_id on ASSISTments/XES3G5M versus exercise-level ucid on Junyi); (iii) a dataset-conditional gate finding on ASSISTments 2012 across four unique learner partitions (five runs, two sharing a split), with occupancy-policy counterfactuals that leave population FAR unchanged (Supplementary Table S4). We do not propose a new KT architecture, a new calibration algorithm, or a new auditing theory, and we do not report a classroom intervention. This is an educational-technology measurement contribution for threshold-based practice systems, not a new pedagogy. A graph-KT (GKT) model and a CL4KT-style adapter appear only as an exploratory single-fold diagnostic on ASSISTments.",
        "Contributions are conservative: (i) a train-only sparse-concept evaluation protocol combining learner-based and temporal views, explicit cold-start definitions, occupancy-aware reporting, and L1–L7 leakage control; (ii) per-stratum calibration diagnostics using ECE, Brier decomposition, and reliability diagrams, showing dataset-dependent rather than universal sparse-calibration vulnerability; (iii) a reproducibility artifact with frozen scripts and configuration and a one-command rebuild of the diagnostic tables from frozen summaries. We do not propose a new KT architecture, a new calibration algorithm, or a classroom intervention. Graph and contrastive KT models are out of scope and are not scored.",
    ),
    (
        "B. Graph and self-supervised KT",
        "B. Evaluation rigor and leakage in KT",
    ),
    (
        "Graph-based KT (GKT) represents knowledge components as nodes of a graph and updates proficiency with a graph neural network [8]. Contrastive learning for KT (CL4KT) trains representations from augmented views of learning histories [9]. These models are related architectures, not the contribution of the present paper. They are included later only as an exploratory single-fold diagnostic on ASSISTments 2012, not as a proposed method and not as a state-of-the-art comparison.",
        "Standard KT benchmarks emphasize overall AUC after a fixed preprocessing recipe [5]. pyKT documents that evaluation choices can inflate discrimination by about 8–13 percent. This paper extends that hygiene to train-only frequency strata and calibration: every bucket, difficulty descriptor, and reporting threshold is fit on the training fold (or selected on validation) and then applied to test. Graph-based and contrastive KT models are related literature only; they are not trained or scored here and are reserved for later studies.",
    ),
    (
        "On ASSISTments fold 0 only we also score a train-only GKT graph [8] (pyKT GKT; Adam, learning rate 1e−3, batch size 16, maximum length 100, hidden and embedding size 32, dropout 0.5, 20 epochs, patience 4, selection by validation AUC) and a CL4KT protocol adapter [9] (contrastive views on training sequences; not an official CL4KT checkpoint; Adam, learning rate 1e−3, batch size 64, maximum length 100, hidden size 64, 4 heads, 2 layers, dropout 0.2, 20 epochs, patience 6, selection by validation AUC).",
        "Graph-based and contrastive KT models are not trained in this study. Table 9 records a seven-channel leakage audit (L1–L7) for the evaluation pipeline on each dataset. Classical BKT is discussed as related work [1] but is not a scored baseline; IRT 1PL is the low-cost classical reference. T-KT is a local Transformer implementation, not published SimpleKT [4].",
    ),
    (
        "H. Simulated decision gate",
        "H. Simulated decision-error probe (not the primary claim)",
    ),
    (
        "Table 5 is a simulated gate at τ=0.7 on ASSISTments 2012 fold 0 (seed 42), not a classroom trial. For T-KT, false-advance rate (FAR) is 0.196 [0.186, 0.208] on dense KCs (N=528,018; Nadvance=284,326; Nincorrect=158,623; E[FAR]=0.113; Excess FAR=0.083; Miss=0.352) and 0.268 [0.202, 0.337] on sparse KCs (N=444, Limited; Nadvance=235; Nincorrect=197; E[FAR]=0.050; Excess FAR=0.218; Miss=0.320). ΔFAR=+0.072. DKT FAR is 0.200 [0.190, 0.211] dense and 0.296 [0.221, 0.383] sparse. Sparse T-KT FAR uses Limited occupancy (Nadvance=235). Exploratory GKT/CL4KT numbers are in subsection E, not in Table 5.",
        "Table 5 is a simulated gate at τ=0.7 on ASSISTments 2012 fold 0 (seed 42), reported as a decision-error probe rather than contribution (ii). For T-KT, false-advance rate (FAR) is 0.196 [0.186, 0.208] on dense KCs (N=528,018; Nadvance=284,326; Nincorrect=158,623; E[FAR]=0.113; Excess FAR=0.083; Miss=0.352) and 0.268 [0.202, 0.337] on sparse KCs (N=444, Limited; Nadvance=235; Nincorrect=197; E[FAR]=0.050; Excess FAR=0.218; Miss=0.320). ΔFAR=+0.072. DKT FAR is 0.200 [0.190, 0.211] dense and 0.296 [0.221, 0.383] sparse. Sparse T-KT FAR uses Limited occupancy (Nadvance=235). Graph and contrastive models are not scored.",
    ),
    (
        "E. Exploratory GKT/CL4KT result",
        "E. Cold-start concept feasibility",
    ),
    (
        "GKT (a train-only graph) and a CL4KT protocol adapter are scored only on ASSISTments 2012 fold 0 (seed 42). They are exploratory, single-fold, ASSISTments-only instantiations: not a state-of-the-art comparison, not a proposed method, and not an official CL4KT checkpoint. On the same seed-42 gate (not shown in Table 5), GKT FAR is 0.205 [0.194, 0.217] dense versus 0.220 [0.149, 0.295] sparse (ΔFAR=+0.015; 95% CI [−0.054, 0.092] includes 0). The CL4KT adapter FAR is 0.185 [0.176, 0.194] versus 0.240 [0.159, 0.330] (ΔFAR 95% CI [−0.018, 0.142] includes 0). A CI that includes 0 is not evidence that either architecture is safer in production.",
        "Strict cold-start is freq_train(c)=0; limited-train cold-start is 0<freq_train(c)<20. Table 11 reports feasibility, not a production claim. On ASSISTments 2012 the strict bin has about four test events (Insufficient) and the very-sparse bin about 17 events (Insufficient), so ECE/AUC on those slices are descriptive only. Junyi’s learner-based ucid tagging yields no estimable sparse or strict-cold-start stratum. XES3G5M very-sparse support is larger (N≈114) but still not a high-N finding. Where the strict slice is too small, the protocol records the empty or Insufficient bucket instead of borrowing a recommender user/item cold-start definition.",
    ),
    (
        "TSCDA is four checks before using one global KT probability threshold for remediation or advancement: (i) inspect KC-frequency occupancy on the training fold actually used, including empty buckets; (ii) evaluate per-stratum calibration separately from AUC; (iii) evaluate threshold-error metrics (FAR, expected FAR, Excess FAR, and Miss) with their denominators; (iv) ensure sufficient sample support before treating a slice as actionable. R/L/I flags are descriptive occupancy labels, not inferential guarantees. This study does not apply temperature scaling [11] or Platt scaling [26]; those maps would be a separate post-hoc experiment. A platform can log Nadvance and FAR by train-only KC stratum from operational traces without an RCT. Those checks are not a validated classroom policy. On the same seed-42 T-KT scores, three simulated policies leave population FAR at 0.197 (Supplementary Table S4): global τ=0.7; advance only on Reliable occupancy; and τ=0.8 on the sparse stratum. Sparse Nadvance is 235, 0, and 218. Population FAR therefore hides the occupancy design; TSCDA reports the slice. A population AUC win, or an exploratory GKT/CL4KT run, is not by itself evidence that a global gate is safer on the tail.",
        "The protocol is four checks before using one global KT probability threshold: (i) inspect KC-frequency occupancy on the training fold actually used, including empty buckets; (ii) evaluate per-stratum ECE and Brier separately from AUC; (iii) if a locked gate is simulated, report FAR with its denominators (Tables 5–6; Supplementary Tables S3–S4); (iv) ensure sufficient sample support before treating a slice as actionable. R/L/I flags are descriptive occupancy labels, not inferential guarantees. This study does not apply temperature scaling [11] or Platt scaling [26]. A population AUC win is not by itself evidence that probabilities are reliable on the tail.",
    ),
    (
        "Next-response correctness is not latent mastery: y=0 is an incorrect next attempt, not a latent-skill diagnosis. The threshold gate is simulated at a locked τ and is not a classroom policy. There is no classroom RCT. GKT and the CL4KT adapter are exploratory, single-fold, ASSISTments-only instantiations, not a state-of-the-art comparison. Temporal evaluation uses a single corrected cutoff (seed 42), not a multi-cutoff variance estimate. Main multi-run summaries use only four unique learner partitions (seeds 2025 and 2026 share a split). R/L/I are descriptive support flags. ECE depends on binning. Three datasets cannot establish a universal diagnostic law. Local IRT, DKT, and T-KT are not byte-identical to a recovered pyKT training image, so Table 2 Version/commit remains NOT RECOVERED; reproduction is from code_for_review_anonymous.zip and the recovered hyperparameters, not from the original training commit.",
        "Next-response correctness is not latent mastery: y=0 is an incorrect next attempt, not a latent-skill diagnosis. The threshold gate is a simulated probe, not a classroom policy. There is no classroom RCT. We do not train graph, contrastive, or self-supervised KT models. BKT is not a scored baseline; IRT is the classical reference (pyBKT 1.4.1 degenerated on ASSISTments seed 42; Supplementary Table S9). T-KT is a local Transformer implementation, not published SimpleKT [4]; official SimpleKT is scored on ASSISTments AUC/ECE only. Temporal evaluation uses a single corrected cutoff (seed 42). Seeds 2025 and 2026 share a split. R/L/I are descriptive support flags. ECE depends on binning. Three datasets cannot establish a universal diagnostic law. The reported diagnostics are reproduced from the frozen review artifact (code_for_review_anonymous.zip) and configuration; locked tables rebuild from frozen summaries (scripts/rebuild_locked_tables.sh).",
    ),
    (
        "Under the evaluated conditions, KT probabilities that look adequate on aggregate AUC can still be poorly calibrated on sparse KCs in some dataset-model settings, and a global threshold can then raise the false-advance rate on those skills. That association appears for T-KT on ASSISTments 2012 (Limited sparse support). It is not reported for Junyi’s empty learner-based sparse bucket, and XES3G5M T-KT ECE remains essentially flat. TSCDA occupancy, calibration, and threshold-error checks are therefore conditionally useful for educational-technology gates. This paper is a simulated decision-error audit, not a new KT model and not a classroom intervention.",
        "Under the evaluated conditions, KT probabilities that look adequate on aggregate AUC can still be poorly calibrated on sparse KCs in some dataset-model settings. That ECE association appears for T-KT on ASSISTments 2012 (Limited sparse support). It is not reported for Junyi’s empty learner-based sparse bucket, and XES3G5M T-KT ECE remains essentially flat. The train-only strata, leakage checklist, and per-stratum ECE/Brier diagnostics are therefore conditionally useful. Graph and contrastive KT, and any path or distillation module, are left to later studies. This paper is a protocol and calibration diagnostic, not a new KT model and not a classroom intervention.",
    ),
    (
        "TSCDA then applies four steps on each fold: (1) assign those train-only strata, including f=0; (2) flag occupancy R/L/I; (3) report ECE/Brier with those flags; (4) apply one locked τ and report FAR with Nadvance (Section III.E–H).",
        "The protocol then applies four steps on each fold: (1) assign those train-only strata, including f=0; (2) flag occupancy R/L/I; (3) report ECE/Brier with those flags; (4) optionally apply one locked τ and report FAR with Nadvance as a probe (Section III.E–H).",
    ),
    (
        "so TSCDA can refuse a sparse claim when tagging is finer than a skill.",
        "so the protocol can refuse a sparse claim when tagging is finer than a skill.",
    ),
]


def set_para_text(p, new: str) -> None:
    if not p.runs:
        p.add_run(new)
        return
    p.runs[0].text = new
    for r in p.runs[1:]:
        r.text = ""


def patch(path: Path) -> None:
    d = Document(str(path))
    hits = 0
    full_before = "\n".join(p.text for p in d.paragraphs)
    for old, new in REPLACEMENTS:
        n = 0
        for p in d.paragraphs:
            if old not in p.text:
                continue
            set_para_text(p, p.text.replace(old, new, 1))
            n += 1
            hits += 1
        if n != 1:
            raise SystemExit(f"{path.name}: hits={n} for {old[:60]!r}")
    full = "\n".join(p.text for p in d.paragraphs)
    leftover = [p.text[:80] for p in d.paragraphs if "TSCDA" in p.text]
    if leftover:
        raise SystemExit(f"{path.name}: TSCDA remains {leftover}")
    if "E. Exploratory GKT" in full or "CL4KT protocol adapter are scored" in full:
        raise SystemExit(f"{path.name}: exploratory GKT/CL4KT remains")
    if "TSCDA" in full_before and "TSCDA" not in full:
        pass
    d.save(str(path))
    print("patched", path.name, "n=", hits)


def main() -> None:
    sys.stdout.reconfigure(encoding="utf-8")
    for name in ("Reproducible Sparse-Concept and Calibration Diagnostics for Knowledge Tracing.docx", "Reproducible Sparse-Concept and Calibration Diagnostics for Knowledge Tracing_blind.docx"):
        patch(HERE / "manuscript" / name)


if __name__ == "__main__":
    main()

#!/usr/bin/env python3
"""D: V.B — EdTech gates consume p, not AUC; protocol is evaluation, not policy."""
from __future__ import annotations

import shutil
import sys
from pathlib import Path

from docx import Document

HERE = Path(__file__).resolve().parent
FULL = HERE / "manuscript" / "Reproducible Sparse-Concept and Calibration Diagnostics for Knowledge Tracing.docx"
BAK = HERE / "manuscript" / "Reproducible Sparse-Concept and Calibration Diagnostics for Knowledge Tracing.docx.bak_pre_p0_d"


def set_para_text(p, new: str) -> None:
    if not p.runs:
        p.add_run(new)
        return
    p.runs[0].text = new
    for r in p.runs[1:]:
        r.text = ""


def main() -> None:
    sys.stdout.reconfigure(encoding="utf-8")
    if not BAK.exists():
        shutil.copy2(FULL, BAK)
    old = "The protocol is four checks before using one global KT probability threshold:"
    added = (
        "Educational-technology gates that skip, remediate, or advance practice "
        "consume a predicted probability, not a population AUC. A ranking win can "
        "therefore hide miscalibration on rarely practiced KCs. This subsection "
        "states an evaluation protocol for that probability-to-decision chain; it "
        "is not a classroom policy and not a recommended production threshold. "
    )
    d = Document(str(FULL))
    n = 0
    for p in d.paragraphs:
        if old not in p.text:
            continue
        set_para_text(p, p.text.replace(old, added + old, 1))
        n += 1
        break
    if n != 1:
        raise SystemExit(f"V.B hits={n}")
    full = "\n".join(p.text for p in d.paragraphs)
    if "not a recommended production threshold" not in full:
        raise SystemExit("D sentence missing")
    if "classroom policy" not in full:
        raise SystemExit("policy disclaimer missing")
    if "0.1136" not in full or "0.2280" not in full:
        raise SystemExit("T-KT ECE lock missing")
    if "TSCDA" in full:
        raise SystemExit("TSCDA reappeared")
    d.save(str(FULL))
    print("D patched")


if __name__ == "__main__":
    main()

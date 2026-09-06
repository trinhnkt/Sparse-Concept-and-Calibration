#!/usr/bin/env python3
"""Read-only format/policy audit of the living IJIET pack."""
from __future__ import annotations

import sys
import zipfile
from pathlib import Path

import fitz
from docx import Document
from docx.oxml.ns import qn

HERE = Path(__file__).resolve().parent.parent
sys.stdout.reconfigure(encoding="utf-8")


def pdf(rel: str):
    d = fitz.open(str(HERE / rel))
    pages = [p.get_text("text") for p in d]
    return d, "\n".join(pages), d.page_count, dict(d.metadata or {}), pages


def main() -> None:
    full_doc, full, fp, fm, fpages = pdf("output/Reproducible Sparse-Concept and Calibration Diagnostics for Knowledge Tracing.pdf")
    blind_doc, blind, bp, bm, _ = pdf("output/Reproducible Sparse-Concept and Calibration Diagnostics for Knowledge Tracing_blind.pdf")
    print("pages", fp, bp)
    print("named meta", fm)
    print("blind meta", bm)

    # dates
    print("--- DATES ---")
    print("Month date, 2026", "Month date, 2026" in full)
    print("September 1", "September 1" in full)
    print("10.18178", "10.18178" in full)
    i = full.find("Manuscript received")
    print(repr(full[i : i + 140].replace("\n", " ")))

    # ethics
    print("--- ETHICS ---")
    i = full.find("Ethical Statement")
    j = full.find("Data and Code Availability")
    eth = full[i:j]
    print(eth)
    print("not under consideration in ethics", "not under consideration" in eth)
    print("not under consideration in named body", "not under consideration" in full)

    # AI
    print("--- AI ---")
    i = full.find("Generative AI Statement")
    j = full.find("REFERENCES") if "REFERENCES" in full[i:] else i + 800
    # find next heading after AI
    k = full.find("REFERENCES", i)
    print(full[i:k][:900])

    # page 3
    print("--- PAGE3 ---")
    print(fpages[2][:1800] if fp >= 3 else "no p3")
    print("p3 chars", len(fpages[2]) if fp >= 3 else 0)

    # tables
    print("--- TABLES ---")
    wd = Document(str(HERE / "manuscript" / "Reproducible Sparse-Concept and Calibration Diagnostics for Knowledge Tracing.docx"))
    for i, tbl in enumerate(wd.tables, 1):
        h = " ".join(tbl.cell(0, 0).text.split())[:70]
        print(f"T{i} {len(tbl.rows)}x{len(tbl.columns)} h={h!r}")

    body = "\n".join(p.text for p in wd.paragraphs)
    for cap in [
        "Table 8. Empirical conditions",
        "Table 8. Cold-start",
        "Table 8. Within-KC",
        "Table 7. Cold-start",
        "Table 7. Empirical",
        "Fig. 3.",
    ]:
        print(cap, cap in body)

    # figure sizes in Word
    print("--- FIGS WORD ---")
    for i, p in enumerate(wd.paragraphs):
        t = p.text.strip()
        if t.startswith("Fig."):
            print("cap", i, t[:120])
    ns = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
    # inline shapes via oxml
    n_blips = 0
    for el in wd.element.iter():
        if el.tag.endswith("}ext") and el.get("cx"):
            cx, cy = el.get("cx"), el.get("cy")
            if cx and int(cx) > 1_000_000:
                print("ext EMU", cx, cy, "pt", round(int(cx) / 12700, 1), round(int(cy) / 12700, 1))
                n_blips += 1
    print("large ext", n_blips)

    # PDF figure bboxes page by page
    print("--- FIGS PDF ---")
    for pi, page in enumerate(full_doc, 1):
        imgs = page.get_images(full=True)
        if imgs:
            print(f"p{pi} images", len(imgs), [(im[2], im[3]) for im in imgs])
        for block in page.get_text("dict")["blocks"]:
            if block.get("type") == 1:
                b = block["bbox"]
                print(f"p{pi} imgbbox", [round(x, 1) for x in b], "w", round(b[2] - b[0], 1), "h", round(b[3] - b[1], 1))

    # zip identity
    print("--- ZIP ---")
    zpath = HERE / "output" / "OJS_UPLOAD" / "code_for_review_anonymous.zip"
    print("exists", zpath.exists(), "bytes", zpath.stat().st_size if zpath.exists() else 0)
    needles = [
        "JEDM",
        "Khanh",
        "trinhnkt",
        "haunv@",
        "utehy",
        "Hung Yen",
        "github.com/trinhnkt",
        "Van-Hau",
        "ORCID",
    ]
    with zipfile.ZipFile(zpath) as zh:
        names = zh.namelist()
        print("nfiles", len(names))
        hits = []
        for n in names:
            raw = zh.read(n)
            try:
                txt = raw.decode("utf-8", "replace")
            except Exception:
                txt = ""
            low = (n + "\n" + txt).lower()
            for nd in needles:
                if nd.lower() in low:
                    hits.append((nd, n))
        print("identity hits", hits[:40], "count", len(hits))

    print("--- BLIND TOKENS ---")
    for nd in ["Khanh-Trinh", "trinhnkt", "Hung Yen", "haunv@", "github.com"]:
        print(nd, nd.lower() in blind.lower() or nd in blind)

    full_doc.close()
    blind_doc.close()


if __name__ == "__main__":
    main()

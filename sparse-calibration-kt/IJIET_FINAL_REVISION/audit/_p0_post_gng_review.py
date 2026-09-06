#!/usr/bin/env python3
"""Post-fix independent review of living named+blind PDFs. Read-only."""
from __future__ import annotations

import sys
from pathlib import Path

import fitz
from docx import Document

HERE = Path(__file__).resolve().parent.parent
OUT = HERE / "audit" / "_p0_post_gng_review.txt"
sys.stdout.reconfigure(encoding="utf-8")


def pdf_text(p: Path) -> tuple[str, int]:
    d = fitz.open(str(p))
    return "\n".join(x.get_text("text") for x in d), d.page_count


def fig_box(pdf: Path, page: int) -> tuple[float, float, int, int]:
    d = fitz.open(str(pdf))
    info = d[page].get_image_info()[0]
    b = info["bbox"]
    return b[2] - b[0], b[3] - b[1], info["width"], info["height"]


def main() -> None:
    blind = HERE / "output" / "Reproducible Sparse-Concept and Calibration Diagnostics for Knowledge Tracing_blind.pdf"
    full = HERE / "output" / "Reproducible Sparse-Concept and Calibration Diagnostics for Knowledge Tracing.pdf"
    docx = HERE / "manuscript" / "Reproducible Sparse-Concept and Calibration Diagnostics for Knowledge Tracing.docx"
    bt, bp = pdf_text(blind)
    ft, fp = pdf_text(full)
    wd = Document(str(docx))
    wt = "\n".join(p.text or "" for p in wd.paragraphs)
    t8 = ""
    for table in wd.tables:
        for row in table.rows:
            a = row.cells[0].text.strip()
            if "tail mass" in a.lower() or "Sparse mass" in a:
                t8 = a
    w1, h1, nw, nh = fig_box(blind, 2)
    w3, h3, n3w, n3h = fig_box(blind, 5)
    lines = [
        f"blind_pages={bp} named_pages={fp}",
        f"fig1_display={w1:.1f}x{h1:.1f} native={nw}x{nh} aspect={w1/h1:.3f} png={nw/nh:.3f}",
        f"fig3_display={w3:.1f}x{h3:.1f}",
        f"t8_row={t8!r}",
    ]
    checks = {
        "pages_9_9": bp == 9 and fp == 9,
        "ece_1136": "0.1136" in bt and "0.2280" in bt,
        "far_196_268": "0.196" in bt and "0.268" in bt,
        "dfar_056": "0.056" in bt,
        "xes_ece": "0.1176" in bt and "0.1129" in bt and "0.1254" in bt,
        "xes_865": "865" in bt and "6,413,353" in bt,
        "xes_114_L": "very-sparse 114 L" in bt.replace("\n", " ") or "114 L 0.184" in bt.replace("\n", " "),
        "no_three_alt": "Three alternative" not in bt and "Three alternative" not in wt,
        "has_two_grids_intro": "Two alternative train-only frequency cut grids" in bt,
        "has_two_grids_results": "two alternative train-only cut grids" in bt,
        "no_competitive": "competitive AUC" not in bt and "competitive AUC" not in wt,
        "has_agg_disc": "aggregate discrimination that does not reveal" in bt,
        "no_success_claims": "Success claims require" not in bt and "Success claims require" not in wt,
        "has_substantive": "Substantive stratum-level interpretations" in bt,
        "no_sparse_mass": "sparse mass" not in bt.lower() and "Sparse mass" not in wt,
        "has_tail_mass": "low-frequency tail mass" in bt.lower(),
        "t8_cell": t8 == "Low-frequency tail mass (E1)",
        "c12_abstract": "learner-based primary reporting with a complementary temporal split" in bt,
        "c12_contrib": "learner-based primary reporting and a complementary temporal split" in bt,
        "fig1_aspect_ok": abs((w1 / h1) - (809 / 268)) < 0.05,
        "fig1_height_ok": h1 > 150,
        "no_tscda": "TSCDA" not in bt,
        "no_chatgpt": "ChatGPT" not in bt and "GPT-5.6" not in bt,
        "grok_only": "Cursor Grok 4.6" in bt,
        "no_gkt_token": "GKT" not in bt and "CL4KT" not in bt,
        "word_E_heading": any(
            (p.text or "").strip() == "E. Secondary explanatory analysis" for p in wd.paragraphs
        ),
        "word_F_heading": any(
            (p.text or "").strip() == "F. Secondary decision-error probe" for p in wd.paragraphs
        ),
        "named_not_anon": "Khanh-Trinh" in ft,
        "blind_anon": "Khanh-Trinh" not in bt and "Anonymous Authors" in bt,
        "placeholder_dates": "Month date, 2026" in bt,
        "far_not_rq3": "not RQ3" in bt,
        "rebuild_script": "rebuild_locked_tables.sh" in bt,
        "official_simplekt": "0.7700" in bt,
        "junyi_ucid": "ucid" in bt and "granularity-and-estimability" in bt,
    }
    # occupancy L context
    flat = " ".join(bt.split())
    checks["xes_vs_flag"] = "XES3G5M very-sparse 114 L" in flat or "XES3G5M very-sparse 114 L 0.184" in flat
    for k, v in checks.items():
        lines.append(f"{'PASS' if v else 'FAIL'} {k}")
    # leftover risk
    for needle in [
        "Three alternative",
        "competitive AUC",
        "Success claims require",
        "plenty of sparse mass",
        "Sparse mass (E1)",
        "TSCDA",
        "five independent",
    ]:
        lines.append(f"LEFT {needle!r} blind={bt.count(needle)} named={ft.count(needle)} word={wt.count(needle)}")
    OUT.write_text("\n".join(lines), encoding="utf-8")
    print("\n".join(lines))


if __name__ == "__main__":
    main()

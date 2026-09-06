#!/usr/bin/env python3
"""Cross-version consistency review of the living IJIET manuscript."""
from __future__ import annotations

import hashlib
import re
import sys
from collections import Counter
from pathlib import Path

import fitz
from docx import Document

HERE = Path(__file__).resolve().parent.parent
MAN = HERE / "manuscript"
OUT = HERE / "output"
OJS = OUT / "OJS_UPLOAD"
SUB_SRC = HERE.parent / "IJIET_SUBMISSION" / "source"
SUB_OUT = HERE.parent / "IJIET_SUBMISSION" / "output"

LOCKS = {
    "0.1136": "ECE dense",
    "0.2280": "ECE sparse",
    "0.196": "FAR dense",
    "0.268": "FAR sparse",
    "Nadvance=235": "sparse Nadvance",
    "[0.006, 0.138]": "locked C2 CI",
    "0.056": "partition ΔFAR",
    "0.015–0.087": "partition range",
    "0.047": "five-run mean",
    "0.1176": "XES dense ECE",
    "0.1129": "XES medium ECE",
    "0.1254": "XES sparse ECE",
    "code_for_review_anonymous.zip": "A33 zip sentence",
    "partition-level mean": "A31 abstract unit",
    "T-KT and DKT only": "Table 5 caption",
    "not shown in Table 5": "IV.E xref",
    "[\u22120.054, 0.092]": "GKT CI",
    "[\u22120.018, 0.142]": "CL4KT CI",
}
FORBID = [
    "4/4 unique partitions (mean ΔFAR 0.047",
    "GKT/CL4KT remain seed 42 only",
    "Threshold-Based Educational Decisions",
    "10.18178",
    "five independent",
]
TITLE = "Reproducible Sparse-Concept and Calibration Diagnostics for Knowledge Tracing"


def sha(p: Path) -> str:
    h = hashlib.sha256()
    with p.open("rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()[:12]


def pdf(path: Path) -> tuple[str, int]:
    d = fitz.open(str(path))
    t = "\n".join(pg.get_text("text") for pg in d)
    n = d.page_count
    d.close()
    return t, n


def docx_text(path: Path) -> str:
    d = Document(str(path))
    paras = [p.text for p in d.paragraphs]
    cells = []
    for tbl in d.tables:
        for row in tbl.rows:
            cells.append(" | ".join(c.text.strip() for c in row.cells))
    return "\n".join(paras + cells)


def norm_body(t: str) -> str:
    t = t.replace("\u2212", "-").replace("–", "-").replace("—", "-")
    t = re.sub(r"\s+", " ", t)
    return t.strip()


def strip_identity(t: str) -> str:
    drop = [
        r"Khanh-Trinh Nguyen.*?accepted Month date, 2026",
        r"Anonymous Authors",
        r"Affiliation withheld for double-blind review",
        r"Corresponding author withheld",
        r"Email withheld for double-blind review",
        r"Hung Yen University of Technology and Education",
        r"Academy of Military Science and Technology",
        r"trinhnk@utehy\.edu\.vn",
        r"tuanymc@utehy\.edu\.vn",
        r"duongnt@utehy\.edu\.vn",
        r"thanhnc@ioit\.ai\.vn",
        r"haunv@utehy\.edu\.vn",
        r"github\.com/trinhnkt",
        r"https://github\.com/[^\s]+",
        r"Acknowledgment\s+We thank.*?(?=References)",
    ]
    out = t
    for pat in drop:
        out = re.sub(pat, " ", out, flags=re.I | re.S)
    return norm_body(out)


def check_set(label: str, t: str, pages: int | None = None) -> list[str]:
    bad = []
    if pages is not None and pages != 8:
        bad.append(f"{label}: pages={pages} not 8")
    if TITLE not in t:
        bad.append(f"{label}: missing title")
    for needle, name in LOCKS.items():
        if needle not in t and needle.replace("\u2212", "-") not in t:
            # hyphen variants
            alt = needle.replace("\u2212", "-")
            if alt not in t.replace("\u2212", "-"):
                bad.append(f"{label}: missing {name} ({needle})")
    for f in FORBID:
        if f in t:
            bad.append(f"{label}: FORBIDDEN {f!r}")
    if "JEDM" in t or "jedm" in t.lower():
        bad.append(f"{label}: names JEDM")
    return bad


def main() -> None:
    sys.stdout.reconfigure(encoding="utf-8")
    files = {
        "live full docx": MAN / "Reproducible Sparse-Concept and Calibration Diagnostics for Knowledge Tracing.docx",
        "live blind docx": MAN / "Reproducible Sparse-Concept and Calibration Diagnostics for Knowledge Tracing_blind.docx",
        "live full pdf": OUT / "Reproducible Sparse-Concept and Calibration Diagnostics for Knowledge Tracing.pdf",
        "live blind pdf": OUT / "Reproducible Sparse-Concept and Calibration Diagnostics for Knowledge Tracing_blind.pdf",
        "ojs full docx": OJS / "Reproducible Sparse-Concept and Calibration Diagnostics for Knowledge Tracing.docx",
        "ojs blind docx": OJS / "Reproducible Sparse-Concept and Calibration Diagnostics for Knowledge Tracing_blind.docx",
        "ojs full pdf": OJS / "Reproducible Sparse-Concept and Calibration Diagnostics for Knowledge Tracing.pdf",
        "ojs blind pdf": OJS / "Reproducible Sparse-Concept and Calibration Diagnostics for Knowledge Tracing_blind.pdf",
        "ojs supp": OJS / "supplementary.pdf",
        "sub full docx": SUB_SRC / "Reproducible Sparse-Concept and Calibration Diagnostics for Knowledge Tracing.docx",
        "sub blind docx": SUB_SRC / "Reproducible Sparse-Concept and Calibration Diagnostics for Knowledge Tracing_blind.docx",
        "sub full pdf": SUB_OUT / "Reproducible Sparse-Concept and Calibration Diagnostics for Knowledge Tracing.pdf",
        "sub blind pdf": SUB_OUT / "Reproducible Sparse-Concept and Calibration Diagnostics for Knowledge Tracing_blind.pdf",
        "sub supp": SUB_OUT / "supplementary.pdf",
        "live supp": OUT / "supplementary.pdf",
    }
    print("=== FILE HASH ===")
    hashes = {}
    for k, p in files.items():
        if not p.is_file():
            print(f"MISSING {k}: {p}")
            continue
        hashes[k] = sha(p)
        print(f"{hashes[k]}  {p.stat().st_size:8d}  {k}")

    pairs = [
        ("live full docx", "ojs full docx"),
        ("live full docx", "sub full docx"),
        ("live blind docx", "ojs blind docx"),
        ("live blind docx", "sub blind docx"),
        ("live full pdf", "ojs full pdf"),
        ("live full pdf", "sub full pdf"),
        ("live blind pdf", "ojs blind pdf"),
        ("live blind pdf", "sub blind pdf"),
        ("live supp", "ojs supp"),
        ("live supp", "sub supp"),
    ]
    print("\n=== COPY IDENTITY ===")
    for a, b in pairs:
        if a not in hashes or b not in hashes:
            print(f"SKIP {a} vs {b}")
            continue
        ok = hashes[a] == hashes[b]
        print(("MATCH" if ok else "DIFF "), a, "<->", b)

    print("\n=== PDF LOCKS / A31-A34 ===")
    issues: list[str] = []
    texts = {}
    for key in ("live full pdf", "live blind pdf", "ojs full pdf", "ojs blind pdf", "sub full pdf", "sub blind pdf"):
        p = files[key]
        t, n = pdf(p)
        texts[key] = t
        print(f"{key}: {n} pages, {len(t)} chars")
        issues.extend(check_set(key, t, n))

    print("\n=== NAMED vs BLIND (identity stripped) ===")
    named_b = strip_identity(texts["live full pdf"])
    blind_b = strip_identity(texts["live blind pdf"])
    if named_b == blind_b:
        print("MATCH scientific body named PDF <-> blind PDF")
    else:
        # show first mismatch window
        for i, (a, b) in enumerate(zip(named_b, blind_b)):
            if a != b:
                print("DIFF at", i)
                print("NAMED", named_b[max(0, i - 80) : i + 120])
                print("BLIND", blind_b[max(0, i - 80) : i + 120])
                break
        else:
            print("DIFF length", len(named_b), len(blind_b))
            print("NAMED tail", named_b[-200:])
            print("BLIND tail", blind_b[-200:])
        issues.append("named/blind scientific body mismatch")

    print("\n=== WORD vs PDF (full) ===")
    w = docx_text(files["live full docx"])
    for needle in (
        "partition-level mean ΔFAR 0.056",
        "Seed-42 ΔFAR 95% CI (KC-cluster, B=2000): T-KT [0.006, 0.138] (locked C2); DKT [0.019, 0.175].",
        "GKT/CL4KT remain seed 42 only",
        "4/4 unique partitions (mean ΔFAR 0.047",
        "code_for_review_anonymous.zip",
        "not shown in Table 5",
    ):
        in_w = needle in w
        in_p = needle in texts["live full pdf"]
        print(f"  {'OK' if in_w == in_p else 'SPLIT'} Word={in_w} PDF={in_p} | {needle[:70]}")
        if in_w != in_p:
            issues.append(f"Word/PDF split: {needle[:50]}")
        if needle.startswith("GKT/CL4KT remain") or needle.startswith("4/4 unique partitions (mean"):
            if in_w or in_p:
                issues.append(f"stale string still present: {needle[:40]}")

    print("\n=== TABLE 5 NOTE 220 ===")
    t = texts["live full pdf"]
    i = t.find("Seed-42 ΔFAR")
    note = t[i : i + 220] if i >= 0 else ""
    print(note.replace("\n", " "))
    if "GKT [" in note or "CL4KT [" in note:
        issues.append("GKT/CL4KT still under Table 5")

    print("\n=== IV.E CIs ===")
    print("GKT CI", ("[\u22120.054, 0.092]" in t) or ("[-0.054, 0.092]" in t.replace("\u2212", "-")))
    print("CL4 CI", ("[\u22120.018, 0.142]" in t) or ("[-0.018, 0.142]" in t.replace("\u2212", "-")))

    print("\n=== BLIND LEAKS ===")
    bt = texts["live blind pdf"]
    for leak in ("Khanh-Trinh", "Hung Yen", "github.com/trinhnkt", "utehy.edu.vn", "ioit.ai.vn"):
        hit = leak.lower() in bt.lower()
        print(("LEAK" if hit else "clean"), leak)
        if hit:
            issues.append(f"blind leak {leak}")

    print("\n=== SI TITLE ===")
    st, sn = pdf(files["live supp"])
    print("SI pages", sn)
    print(st[:240].replace("\n", " | "))
    if TITLE not in st:
        issues.append("SI title mismatch")

    print("\n=== COVER vs ARTICLE ===")
    cover = (OUT / "cover_letter_ijiet.txt").read_text(encoding="utf-8")
    print("cover JEDM", "JEDM" in cover)
    print("article JEDM", "JEDM" in texts["live full pdf"])
    if "JEDM" not in cover:
        issues.append("cover missing JEDM withdrawal note")

    print("\n=== TABLE 6 FONT (Word) ===")
    d = Document(str(files["live full docx"]))
    for tbl in d.tables:
        if len(tbl.rows) < 2:
            continue
        if tbl.rows[1].cells[0].text.strip() != "T-KT":
            continue
        if "0.056" not in tbl.rows[1].cells[1].text:
            continue
        for r_i, row in enumerate(tbl.rows):
            sizes = []
            for c in row.cells:
                sz = []
                for p in c.paragraphs:
                    for run in p.runs:
                        if run.font.size:
                            sz.append(round(run.font.size.pt, 1))
                sizes.append((c.text.strip()[:24], sz))
            print(f"  r{r_i}", sizes)

    print("\n=== STALE README ===")
    readme = (OJS / "README_SUBMIT.md").read_text(encoding="utf-8")
    print("README first line:", readme.splitlines()[0])
    if "A29" in readme.splitlines()[0]:
        print("NOTE README still says A29 (pack label stale, not article text)")

    print("\n=== ISSUES ===")
    if issues:
        for x in issues:
            print("!", x)
    else:
        print("none")


if __name__ == "__main__":
    main()

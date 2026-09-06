#!/usr/bin/env python3
from __future__ import annotations
import sys
from docx import Document

sys.stdout.reconfigure(encoding="utf-8")
d = Document(r"IJIET_FINAL_REVISION/manuscript/Reproducible Sparse-Concept and Calibration Diagnostics for Knowledge Tracing.docx")
for i in (69, 70, 88, 89):
    print("====", i)
    print(d.paragraphs[i].text)
    print()
body = "\n".join(p.text for p in d.paragraphs)
for s in (
    "Table 7",
    "Table 8",
    "IV.C",
    "IV.D",
    "IV.E",
    "Section IV",
):
    print(s, body.count(s))

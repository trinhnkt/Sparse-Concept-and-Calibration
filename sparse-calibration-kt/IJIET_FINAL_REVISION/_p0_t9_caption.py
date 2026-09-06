#!/usr/bin/env python3
"""Table 9 caption: XES very-sparse is the same masked series as Table 5."""
from __future__ import annotations

import shutil
import sys
from pathlib import Path

from docx import Document

HERE = Path(__file__).resolve().parent
FULL = HERE / "manuscript" / "Reproducible Sparse-Concept and Calibration Diagnostics for Knowledge Tracing.docx"
BAK = HERE / "manuscript" / "Reproducible Sparse-Concept and Calibration Diagnostics for Knowledge Tracing.docx.bak_pre_t9_caption"


def set_para_text(p, new: str) -> None:
    if not p.runs:
        p.add_run(new)
        return
    p.runs[0].text = new
    for r in p.runs[1:]:
        r.text = ""


def main() -> None:
    sys.stdout.reconfigure(encoding="utf-8")
    if not BAK.exists():
        shutil.copy2(FULL, BAK)
    old = (
        "I = Insufficient (N<100). Not a recommender user/item split."
    )
    extra = (
        " XES3G5M very-sparse uses the same masked four-partition series "
        "as Table 5 (N=114)."
    )
    d = Document(str(FULL))
    n = 0
    for p in d.paragraphs:
        if old not in p.text or "Table 9." not in p.text:
            continue
        if "same masked four-partition series" in p.text:
            print("caption already patched")
            return
        set_para_text(p, p.text.replace(old, old + extra, 1))
        n += 1
        break
    if n != 1:
        raise SystemExit(f"caption hits={n}")
    full = "\n".join(p.text for p in d.paragraphs)
    if "same masked four-partition series" not in full:
        raise SystemExit("caption clause missing")
    if "N≈114" not in full:
        raise SystemExit("N≈114 prose missing")
    if "0.1136" not in full or "0.2280" not in full:
        raise SystemExit("ECE lock missing")
    d.save(str(FULL))
    print("t9 caption patched")


if __name__ == "__main__":
    main()

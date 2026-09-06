#!/usr/bin/env python3
"""Pointers for S8 (bucket AUC/ACC) and S9 / BKT fallback. No numeric locks changed."""
from __future__ import annotations

import shutil
import sys
from pathlib import Path

from docx import Document

HERE = Path(__file__).resolve().parent
FULL = HERE / "manuscript" / "Reproducible Sparse-Concept and Calibration Diagnostics for Knowledge Tracing.docx"
BAK = HERE / "manuscript" / "Reproducible Sparse-Concept and Calibration Diagnostics for Knowledge Tracing.docx.bak_pre_p0_partial3"


def set_para_text(p, new: str) -> None:
    if not p.runs:
        p.add_run(new)
        return
    p.runs[0].text = new
    for r in p.runs[1:]:
        r.text = ""


def main() -> None:
    sys.stdout.reconfigure(encoding="utf-8")
    if not BAK.exists():
        shutil.copy2(FULL, BAK)
    pairs = [
        (
            "T-KT is not that checkpoint.",
            "T-KT is not that checkpoint. Stratum-level AUC and ACC for IRT, "
            "DKT, and T-KT are in Supplementary Table S8.",
        ),
        (
            "BKT is not a scored baseline; IRT is the classical reference. "
            "Official SimpleKT [4] is scored on ASSISTments AUC/ECE only.",
            "BKT is not a scored baseline; IRT is the classical reference "
            "(pyBKT 1.4.1 degenerated on ASSISTments seed 42; Supplementary "
            "Table S9). Official SimpleKT [4] is scored on ASSISTments AUC/ECE "
            "only and is not compared to a published ASSISTments 2012 pyKT "
            "cell, which that paper does not report.",
        ),
        (
            "reproduction is from code_for_review_anonymous.zip and the recovered hyperparameters.",
            "reproduction is from code_for_review_anonymous.zip and the recovered "
            "hyperparameters. Locked Table 5/6 cells are rebuilt from frozen "
            "prediction summaries (scripts/rebuild_locked_tables.sh), not by retraining.",
        ),
    ]
    d = Document(str(FULL))
    for old, new in pairs:
        n = 0
        for p in d.paragraphs:
            if old not in p.text:
                continue
            set_para_text(p, p.text.replace(old, new, 1))
            n += 1
            break
        if n != 1:
            raise SystemExit(f"hits={n} for {old[:80]!r}")
    full = "\n".join(p.text for p in d.paragraphs)
    if "Supplementary Table S8" not in full:
        raise SystemExit("S8 pointer missing")
    if "Supplementary Table S9" not in full:
        raise SystemExit("S9 pointer missing")
    if "rebuild_locked_tables.sh" not in full:
        raise SystemExit("rebuild script pointer missing")
    if "0.1136" not in full or "0.2280" not in full:
        raise SystemExit("T-KT ECE lock missing")
    if "TSCDA" in full:
        raise SystemExit("TSCDA reappeared")
    d.save(str(FULL))
    print("partial3 patched")


if __name__ == "__main__":
    main()

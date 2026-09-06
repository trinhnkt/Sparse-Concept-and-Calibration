#!/usr/bin/env python3
"""IV.E + Table 9 caption: 0<f<20 is very-sparse, not limited cold-start."""
from __future__ import annotations

import shutil
import sys
from pathlib import Path

from docx import Document

HERE = Path(__file__).resolve().parent
FULL = HERE / "manuscript" / "Reproducible Sparse-Concept and Calibration Diagnostics for Knowledge Tracing.docx"
BAK = HERE / "manuscript" / "Reproducible Sparse-Concept and Calibration Diagnostics for Knowledge Tracing.docx.bak_pre_p0_cs_label"

OLD_E = (
    "Strict cold-start is freq_train(c)=0; limited-train cold-start is "
    "0<freq_train(c)<20."
)
NEW_E = (
    "Strict cold-start is freq_train(c)=0. The next Table 9 slice is the "
    "protocol very-sparse bucket (0<freq_train(c)<20), not a limited "
    "cold-start (k5/k10) group."
)

OLD_CAP = "Strict: f_train=0. Limited/very-sparse: 0<f_train<20. "
NEW_CAP = (
    "Strict: f_train=0. Very-sparse: 0<f_train<20. L/I are occupancy "
    "flags (N), not k5/k10 groups. "
)


def set_para_text(p, new: str) -> None:
    if not p.runs:
        p.add_run(new)
        return
    p.runs[0].text = new
    for r in p.runs[1:]:
        r.text = ""


def main() -> None:
    sys.stdout.reconfigure(encoding="utf-8")
    if not BAK.exists():
        shutil.copy2(FULL, BAK)
    d = Document(str(FULL))
    hits = {"e": 0, "cap": 0}
    for p in d.paragraphs:
        if OLD_E in p.text:
            set_para_text(p, p.text.replace(OLD_E, NEW_E, 1))
            hits["e"] += 1
        elif OLD_CAP in p.text and "Table 9." in p.text:
            set_para_text(p, p.text.replace(OLD_CAP, NEW_CAP, 1))
            hits["cap"] += 1
    if hits["e"] != 1 or hits["cap"] != 1:
        raise SystemExit(f"hits={hits}")
    paras = "\n".join(p.text for p in d.paragraphs)
    cells = "\n".join(
        c.text for tbl in d.tables for row in tbl.rows for c in row.cells
    )
    full = paras + "\n" + cells
    if "limited-train cold-start is 0<" in paras:
        raise SystemExit("old IV.E label remains")
    if "Limited/very-sparse" in paras:
        raise SystemExit("old Table 9 caption remains")
    if "not a limited cold-start (k5/k10) group" not in paras:
        raise SystemExit("new IV.E missing")
    if "not k5/k10 groups" not in paras:
        raise SystemExit("new caption missing")
    if "0.245" not in cells or "0.178" not in cells:
        raise SystemExit("Assist Table 9 ECE lock missing")
    if "N≈114" not in paras:
        raise SystemExit("XES N≈114 missing")
    if "0.1136" not in full or "0.2280" not in full:
        raise SystemExit("T-KT ECE lock missing")
    d.save(str(FULL))
    print("cs label patched")


if __name__ == "__main__":
    main()

#!/usr/bin/env python3
"""IV.B: one S7 sentence; calibration ordering unchanged (T-KT+DKT)."""
from __future__ import annotations

import shutil
import sys
from pathlib import Path

from docx import Document

HERE = Path(__file__).resolve().parent
FULL = HERE / "manuscript" / "Reproducible Sparse-Concept and Calibration Diagnostics for Knowledge Tracing.docx"
BAK = HERE / "manuscript" / "Reproducible Sparse-Concept and Calibration Diagnostics for Knowledge Tracing.docx.bak_pre_p0_ivb_s7"

OLD = (
    "Two alternative train-only cut grids on the same frozen T-KT scores "
    "leave that dense-to-sparse rise positive (Supplementary Table S7)."
)
NEW = (
    "The main ASSISTments calibration ordering remains qualitatively "
    "unchanged under two alternative train-only frequency cut grids "
    "(Supplementary Table S7)."
)


def set_para_text(p, new: str) -> None:
    if not p.runs:
        p.add_run(new)
        return
    p.runs[0].text = new
    for r in p.runs[1:]:
        r.text = ""


def main() -> None:
    sys.stdout.reconfigure(encoding="utf-8")
    if not BAK.exists():
        shutil.copy2(FULL, BAK)
    d = Document(str(FULL))
    n = 0
    for p in d.paragraphs:
        if OLD not in p.text:
            continue
        set_para_text(p, p.text.replace(OLD, NEW, 1))
        n += 1
    if n != 1:
        raise SystemExit(f"hits={n}")
    full = "\n".join(p.text for p in d.paragraphs)
    if OLD in full:
        raise SystemExit("old IV.B remains")
    if "qualitatively unchanged" not in full:
        raise SystemExit("new IV.B missing")
    if full.count("Supplementary Table S7") < 2:
        raise SystemExit("S7 pointers dropped")
    if "0.1136" not in full or "0.2280" not in full:
        raise SystemExit("T-KT ECE lock missing")
    d.save(str(FULL))
    print("ivb s7 patched")


if __name__ == "__main__":
    main()

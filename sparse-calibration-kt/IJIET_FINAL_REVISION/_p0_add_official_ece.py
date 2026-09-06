#!/usr/bin/env python3
"""Add official SimpleKT ASSISTments AUC/ECE. Do not change T-KT 0.1136/0.2280."""
from __future__ import annotations

import shutil
import sys
from pathlib import Path

import win32com.client as win32
from docx import Document

HERE = Path(__file__).resolve().parent
FULL = HERE / "manuscript" / "Reproducible Sparse-Concept and Calibration Diagnostics for Knowledge Tracing.docx"
BAK = HERE / "manuscript" / "Reproducible Sparse-Concept and Calibration Diagnostics for Knowledge Tracing.docx.bak_pre_official_ece"
WD_SAVE = -1
PM = "\u00b1"


def txt(cell) -> str:
    return cell.Range.Text.replace("\r", " ").replace("\x07", "").strip()


def set_cell(cell, text: str) -> None:
    cell.Range.Text = text


def insert_rows() -> None:
    word = win32.DispatchEx("Word.Application")
    word.Visible = False
    word.DisplayAlerts = 0
    doc = word.Documents.Open(str(FULL))
    log = []
    try:
        t4 = None
        t5 = None
        for t in range(1, doc.Tables.Count + 1):
            tbl = doc.Tables(t)
            if tbl.Rows.Count == 10 and tbl.Columns.Count == 4 and txt(tbl.Cell(1, 3)) == "AUC":
                t4 = tbl
            if tbl.Rows.Count == 9 and tbl.Columns.Count == 5 and txt(tbl.Cell(1, 5)) == "ECE":
                t5 = tbl
        if t4 is None or t5 is None:
            raise SystemExit(f"tables not found t4={t4} t5={t5}")
        if txt(t4.Cell(4, 2)) != "T-KT" or "0.6837" not in txt(t4.Cell(4, 3)):
            raise SystemExit(f"unexpected t4 row4 {txt(t4.Cell(4, 2))} {txt(t4.Cell(4, 3))}")
        if txt(t4.Cell(5, 2)) != "SimpleKT":
            t4.Rows.Add(t4.Rows(5))
            set_cell(t4.Cell(5, 1), "ASSISTments 2012")
            set_cell(t4.Cell(5, 2), "SimpleKT")
            set_cell(t4.Cell(5, 3), f"0.7700{PM}0.0013")
            set_cell(t4.Cell(5, 4), f"0.7522{PM}0.0014")
            log.append("t4 row SimpleKT")
        if "0.1136" not in txt(t5.Cell(2, 5)) or "0.2280" not in txt(t5.Cell(4, 5)):
            raise SystemExit("T-KT ECE lock missing in table 5")
        if txt(t5.Cell(5, 1)) != "Assist. SimpleKT":
            for _ in range(3):
                t5.Rows.Add(t5.Rows(5))
            for r, (st, n, fl, ece) in enumerate(
                (
                    ("dense", "523,971", "R", f"0.0203{PM}0.0035"),
                    ("medium", "5,963", "R", f"0.0262{PM}0.0025"),
                    ("sparse", "415", "L", f"0.0884{PM}0.0187"),
                ),
                start=5,
            ):
                set_cell(t5.Cell(r, 1), "Assist. SimpleKT")
                set_cell(t5.Cell(r, 2), st)
                set_cell(t5.Cell(r, 3), n)
                set_cell(t5.Cell(r, 4), fl)
                set_cell(t5.Cell(r, 5), ece)
            log.append("t5 3 SimpleKT rows")
        if "0.1136" not in txt(t5.Cell(2, 5)) or "0.2280" not in txt(t5.Cell(4, 5)):
            raise SystemExit("T-KT ECE overwritten")
        doc.Save()
        log.append("rows saved")
    finally:
        try:
            doc.Close(WD_SAVE)
        except Exception:
            pass
        word.Quit()
    print("\n".join(log))


def set_para_text(p, new: str) -> None:
    if not p.runs:
        p.add_run(new)
        return
    p.runs[0].text = new
    for r in p.runs[1:]:
        r.text = ""


def patch_prose() -> None:
    pairs = [
        (
            "SimpleKT [4] is related literature only; it is not the T-KT model scored in Section IV.",
            "SimpleKT [4] is not the T-KT model; official SimpleKT is scored on ASSISTments AUC/ECE only.",
        ),
        (
            "not the published SimpleKT architecture [4]; no Section IV cell is a SimpleKT [4] result.",
            "not the published SimpleKT architecture [4]. Official SimpleKT [4] is scored on ASSISTments AUC/ECE only.",
        ),
        (
            "DKT is slightly above T-KT on all three datasets "
            f"(ASSISTments DKT 0.6979{PM}0.0014 versus T-KT 0.6837{PM}0.0025). "
            "In our implementation, unseen learners trigger",
            "DKT is slightly above T-KT on all three datasets "
            f"(ASSISTments DKT 0.6979{PM}0.0014 versus T-KT 0.6837{PM}0.0025). "
            "Official SimpleKT [4] on the same ASSISTments partitions reaches "
            f"0.7700{PM}0.0013 (ACC 0.7522{PM}0.0014); T-KT is not that checkpoint. "
            "In our implementation, unseen learners trigger",
        ),
        (
            "That sparse cell is Limited support, not a high-N finding.",
            "That sparse cell is Limited support, not a high-N finding. "
            f"Official SimpleKT [4] on those partitions has ECE 0.0203{PM}0.0035 "
            f"dense and 0.0884{PM}0.0187 sparse (Limited, N=415); the rise remains, "
            "and the T-KT cells stay locked.",
        ),
        (
            "BKT is not a scored baseline; IRT is the classical reference.",
            "BKT is not a scored baseline; IRT is the classical reference. "
            "Official SimpleKT [4] is scored on ASSISTments AUC/ECE only.",
        ),
        (
            "Table 4. Overall learner-based area under the ROC curve (AUC) and "
            f"accuracy (ACC) (mean{PM}sd over four unique partitions).",
            "Table 4. Overall learner-based area under the ROC curve (AUC) and "
            f"accuracy (ACC) (mean{PM}sd over four unique partitions). Official "
            "SimpleKT [4] is reported on ASSISTments only.",
        ),
        (
            "Table 5. T-KT event-level expected calibration error (ECE) by "
            "train-only frequency stratum.",
            "Table 5. T-KT and official SimpleKT [4] (ASSISTments only) "
            "event-level expected calibration error (ECE) by train-only "
            "frequency stratum.",
        ),
    ]
    d = Document(str(FULL))
    hits = 0
    for old, new in pairs:
        n = 0
        for p in d.paragraphs:
            if old not in p.text:
                continue
            set_para_text(p, p.text.replace(old, new, 1))
            n += 1
            hits += 1
        if n != 1:
            raise SystemExit(f"prose hits={n} for {old[:70]!r}")
    full = "\n".join(p.text for p in d.paragraphs)
    if "0.1136" not in full or "0.2280" not in full:
        raise SystemExit("T-KT ECE lock missing after prose")
    if "no Section IV cell is a SimpleKT" in full:
        raise SystemExit("stale no-SimpleKT-cell sentence remains")
    d.save(str(FULL))
    print(f"prose patched n={hits}")


def main() -> None:
    sys.stdout.reconfigure(encoding="utf-8")
    if not BAK.exists():
        shutil.copy2(FULL, BAK)
    insert_rows()
    patch_prose()


if __name__ == "__main__":
    main()

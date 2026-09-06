#!/usr/bin/env python3
"""Junyi terminology: ucid is an operational identifier, not a pedagogical KC."""
from __future__ import annotations

import shutil
import sys
from pathlib import Path

from docx import Document

HERE = Path(__file__).resolve().parent
FULL = HERE / "manuscript" / "Reproducible Sparse-Concept and Calibration Diagnostics for Knowledge Tracing.docx"
BAK = HERE / "manuscript" / "Reproducible Sparse-Concept and Calibration Diagnostics for Knowledge Tracing.docx.bak_pre_p0_junyi_term"

PAIRS = [
    (
        "Junyi’s exercise-level KC tagging yields no learner-based sparse "
        "stratum (estimability, not a missing table)",
        "under Junyi Academy’s exercise-level operational identifier (ucid), "
        "no learner-based sparse stratum is formed under the pre-specified "
        "thresholds",
    ),
    (
        "Junyi’s learner-based ucid tagging yields no estimable sparse or "
        "strict-cold-start stratum.",
        "Junyi’s learner-based operational identifier (ucid) yields no "
        "estimable sparse or strict-cold-start stratum.",
    ),
    (
        "Junyi’s learner-based sparse bucket is empty (exercise-level tagging) "
        "rather than a ranking collapse.",
        "Junyi’s learner-based sparse bucket is empty (exercise-level "
        "operational identifier, ucid) rather than a ranking collapse.",
    ),
    (
        "Junyi’s learner-based sparse bucket is empty because the operational "
        "KC is an exercise identifier (ucid), so the protocol can refuse a "
        "sparse claim when tagging is finer than a skill.",
        "Junyi’s learner-based sparse bucket is empty because the operational "
        "field is ucid (an exercise-level identifier, not a skill tag), so the "
        "protocol can refuse a sparse claim when the unit is finer than a "
        "skill. Junyi is therefore interpreted primarily as a "
        "granularity-and-estimability case rather than direct evidence about "
        "pedagogical KC sparsity.",
    ),
]


def set_para_text(p, new: str) -> None:
    if not p.runs:
        p.add_run(new)
        return
    p.runs[0].text = new
    for r in p.runs[1:]:
        r.text = ""


def main() -> None:
    sys.stdout.reconfigure(encoding="utf-8")
    if not BAK.exists():
        shutil.copy2(FULL, BAK)
    d = Document(str(FULL))
    hits = {i: 0 for i in range(len(PAIRS))}
    for p in d.paragraphs:
        t = p.text
        for i, (old, new) in enumerate(PAIRS):
            if old in t:
                set_para_text(p, t.replace(old, new, 1))
                hits[i] += 1
                t = p.text
    if any(hits[i] != 1 for i in hits):
        raise SystemExit(f"hits={hits}")
    full = "\n".join(p.text for p in d.paragraphs)
    if "exercise-level KC tagging" in full:
        raise SystemExit("old Abstract KC tagging remains")
    if "ucid tagging" in full:
        raise SystemExit("old IV.E ucid tagging remains")
    if "exercise-level tagging)" in full:
        raise SystemExit("old V.A tagging remains")
    if "operational KC is an exercise identifier" in full:
        raise SystemExit("old V.C operational KC remains")
    if "granularity-and-estimability case" not in full:
        raise SystemExit("V.C sentence missing")
    if "exercise-level operational identifier (ucid)" not in full:
        raise SystemExit("Abstract wording missing")
    if "0.1136" not in full or "0.2280" not in full:
        raise SystemExit("T-KT ECE lock missing")
    if "TSCDA" in full:
        raise SystemExit("TSCDA reappeared")
    d.save(str(FULL))
    print("junyi term patched", hits)


if __name__ == "__main__":
    main()
